from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[3]
OUT_NAME = "Appendix_Two_More_Functionalities_Code.docx"
OUT_PATHS = [
    ROOT / OUT_NAME,
    Path.home() / "OneDrive" / "Desktop" / OUT_NAME,
    Path.home() / "Downloads" / OUT_NAME,
]


def set_run_font(run, name="Times New Roman", size=12, bold=False):
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), name)


def add_heading_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=14, bold=True)


def add_body(doc, text, *, italic=False, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=size)
    run.italic = italic


def add_code(doc, text):
    for line in text.splitlines() or [""]:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=9)


def slice_file(path: Path, start: int, end: int) -> str:
    lines = path.read_text(encoding="utf-8").splitlines()
    return "\n".join(lines[start - 1 : end])


backend = ROOT / "backend"

sections = [
    (
        "1. Create Project",
        "File: backend/projects/views.py  |  API: POST /api/projects/",
        slice_file(backend / "projects" / "views.py", 1, 152),
    ),
    (
        "2. Create Worker",
        "File: backend/accounts/user_views.py  |  API: POST /api/users/",
        (backend / "accounts" / "user_views.py").read_text(encoding="utf-8"),
    ),
    (
        "3. Create Milestone",
        "File: backend/projects/views.py  |  API: POST /api/projects/<id>/milestones/",
        slice_file(backend / "projects" / "views.py", 245, 375),
    ),
    (
        "4. Assign Task",
        "File: backend/projects/execution_views.py  |  API: POST /api/projects/<id>/milestones/<id>/tasks/<id>/assignments/",
        slice_file(backend / "projects" / "execution_views.py", 39, 59)
        + "\n\n"
        + slice_file(backend / "projects" / "execution_views.py", 511, 574)
        + "\n\n"
        + slice_file(backend / "projects" / "execution_serializers.py", 144, 163)
        + "\n\n"
        + slice_file(backend / "projects" / "execution_services.py", 837, 888),
    ),
    (
        "5. Worker Clock In and Clock Out",
        "File: backend/attendance/views.py and backend/attendance/services.py  |  API: POST /api/attendance/clock-in/ and POST /api/attendance/clock-out/",
        (backend / "attendance" / "views.py").read_text(encoding="utf-8")
        + "\n\n"
        + slice_file(backend / "attendance" / "services.py", 210, 298),
    ),
    (
        "6. Create Task",
        "File: backend/projects/execution_views.py  |  API: POST /api/projects/<id>/milestones/<id>/tasks/",
        slice_file(backend / "projects" / "execution_views.py", 227, 362)
        + "\n\n"
        + slice_file(backend / "projects" / "execution_serializers.py", 80, 141)
        + "\n\n"
        + slice_file(backend / "projects" / "execution_services.py", 652, 686),
    ),
    (
        "7. Create Project Document",
        "File: backend/projects/views.py  |  API: POST /api/projects/<id>/documents/",
        slice_file(backend / "projects" / "views.py", 378, 510)
        + "\n\n"
        + slice_file(backend / "projects" / "serializers.py", 575, 585)
        + "\n\n"
        + slice_file(backend / "projects" / "services.py", 682, 691),
    ),
]

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)

title = doc.add_paragraph()
run = title.add_run("APPENDIX")
set_run_font(run, size=16, bold=True)

sub = doc.add_paragraph()
run = sub.add_run("Source Code of Major Functionalities — FORTESITE")
set_run_font(run, size=13, bold=True)

add_body(
    doc,
    "This appendix presents two additional functionalities of FORTESITE: Create Task and Create Project Document. Each listing includes authentication or role checking before data is saved.",
)

for heading, location, code in sections[-2:]:
    add_heading_line(doc, heading)
    add_body(doc, location, italic=True, size=10)
    add_code(doc, code)

saved = []
for out in OUT_PATHS:
    try:
        doc.save(out)
        saved.append(str(out))
    except OSError as exc:
        print(f"skip {out}: {exc}")
print("\n".join(saved))
