from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, Twips

ROOT = Path(__file__).resolve().parents[3]
OUT_PATHS = [
    ROOT / "Login_Test_Case_1.docx",
    Path.home() / "OneDrive" / "Desktop" / "Login_Test_Case_1.docx",
    Path.home() / "OneDrive" / "Desktop" / "Login_Test_Case_1_Amal.docx",
    Path.home() / "Downloads" / "Login_Test_Case_1.docx",
    Path.home() / "Documents" / "Login_Test_Case_1.docx",
]

FONT = "Times New Roman"
# A4 content width with 2.54 cm margins = 15.92 cm (matches printed report table)
TABLE_CM = 15.92
COL_CM = [1.35, 2.75, 2.45, 2.75, 2.75, 3.87]


def cm_dxa(cm: float) -> int:
    return int(round(cm * 567))


def set_font(run, *, size=11, bold=False, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = FONT
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)


def tighten(p, *, center=False, size=11):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag in (qn("w:tcBorders"), qn("w:tcMar")):
            tcPr.remove(child)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "60"), ("left", "80"), ("bottom", "60"), ("right", "80")):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), val)
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def set_cell_width(cell, cm: float):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcW"):
            tcPr.remove(child)
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(cm_dxa(cm)))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)
    cell.width = Cm(cm)


def write(cell, text, *, size=11, bold=False, center=False, vcenter=True):
    cell.text = ""
    lines = str(text).split("\n") if text is not None else [""]
    if not lines:
        lines = [""]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        tighten(p, center=center, size=size)
        run = p.add_run(line)
        set_font(run, size=size, bold=bold)
    set_cell_border(cell)
    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER if vcenter else WD_CELL_VERTICAL_ALIGNMENT.TOP
    )


def merge(table, r1, c1, r2, c2, text, **kwargs):
    a = table.cell(r1, c1)
    b = table.cell(r2, c2)
    a.merge(b)
    write(a, text, **kwargs)
    return a


def set_row_height(row, cm: float):
    row.height = Cm(cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def set_fixed_layout(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    for child in list(tblPr):
        if child.tag in (qn("w:tblW"), qn("w:tblLayout"), qn("w:tblBorders"), qn("w:tblGrid")):
            tblPr.remove(child)

    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(cm_dxa(TABLE_CM)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    jc = tblPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblPr.append(jc)
    jc.set(qn("w:val"), "center")

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)

    for child in list(tbl):
        if child.tag == qn("w:tblGrid"):
            tbl.remove(child)
    grid = OxmlElement("w:tblGrid")
    for w in COL_CM:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(cm_dxa(w)))
        grid.append(gc)
    tbl.insert(1, grid)


def add_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    tighten(hp)
    hp.paragraph_format.tab_stops.add_tab_stop(
        Cm(TABLE_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
    )
    r1 = hp.add_run("FORTESITE")
    set_font(r1, size=11)
    r2 = hp.add_run("\t")
    set_font(r2, size=11)
    r3 = hp.add_run("1")
    set_font(r3, size=11)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    tighten(fp)
    fp.paragraph_format.tab_stops.add_tab_stop(
        Cm(TABLE_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
    )
    left = fp.add_run("Amal Jyothi College of Engineering Autonomous, Kanjirappally")
    set_font(left, size=9, italic=True)
    fp.add_run("\t")
    right = fp.add_run("Department of Computer Applications")
    set_font(right, size=9, italic=True)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.54)
section.right_margin = Cm(2.54)
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.header_distance = Cm(1.25)
section.footer_distance = Cm(1.25)
add_header_footer(section)

title = doc.add_paragraph()
tighten(title)
tr = title.add_run("Test report")
set_font(tr, size=14, bold=True)
title.paragraph_format.space_after = Pt(8)
# remove the blank default paragraph if present
body = doc.element.body
for child in list(body):
    if child.tag == qn("w:p") and child is not title._p and not (child.text or "").strip():
        if child.find(qn("w:r")) is None:
            body.remove(child)
            break

table = doc.add_table(rows=15, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.allow_autofit = False
set_fixed_layout(table)

for row in table.rows:
    for i, cell in enumerate(row.cells):
        set_cell_width(cell, COL_CM[i])

# Same bold pattern as the sample: left labels/titles bold, right names/dates regular
merge(table, 0, 0, 0, 5, "Test Case 1", bold=True)
merge(table, 1, 0, 1, 5, "Project Name: FORTESITE", bold=True)
merge(table, 2, 0, 2, 5, "Login Test Case", bold=True)

merge(table, 3, 0, 3, 2, "Test Case ID: Test_1", bold=True)
merge(table, 3, 3, 3, 5, "Test Designed By: Liya Mary Biju", bold=False)

merge(table, 4, 0, 4, 2, "Test Priority(Low/Medium/High): High", bold=True)
merge(table, 4, 3, 4, 5, "Test Designed Date: 23/10/2026", bold=False)

merge(table, 5, 0, 5, 2, "Module Name: Login page", bold=True)
merge(table, 5, 3, 5, 5, "Test Executed By : Amal K Jose", bold=False)

merge(
    table,
    6,
    0,
    6,
    2,
    "Test Title : verify login with valid email and password.",
    bold=True,
)
merge(table, 6, 3, 6, 5, "Test Execution Date: 23/10/2026", bold=False)

merge(table, 7, 0, 7, 2, "Description:", bold=True)
merge(table, 7, 3, 7, 5, "User has valid email and password.", bold=False)

merge(table, 8, 0, 8, 5, "Pre-Condition :User has valid email and password", bold=True)

headers = [
    "Step",
    "Test Step",
    "Test Data",
    "Expected Result",
    "Actual Result",
    "Status(Pass/\nFail)",
]
for i, h in enumerate(headers):
    write(table.cell(9, i), h, bold=True)

write(table.cell(10, 0), "1")
write(table.cell(10, 1), "Navigate to the login page")
write(table.cell(10, 2), "http://localhost:5173/login")
write(table.cell(10, 3), "Display login page")
write(table.cell(10, 4), "Login page displayed")
write(table.cell(10, 5), "pass")

write(table.cell(11, 0), "2")
write(table.cell(11, 1), "Valid email entered")
write(table.cell(11, 2), "Email:\nadmin@example.com")

write(table.cell(12, 0), "3")
write(table.cell(12, 1), "Valid password entered")
write(table.cell(12, 2), "Password:\n123")

write(table.cell(13, 0), "4")
write(table.cell(13, 1), "Click on Sign In button")
write(table.cell(13, 2), "")

merge(table, 11, 3, 13, 3, "Admin should enter the main page")
merge(table, 11, 4, 13, 4, "Admin entered to the main page")
merge(table, 11, 5, 13, 5, "pass")

merge(
    table,
    14,
    0,
    14,
    5,
    "Post-Condition: Email and password is checked with database values for successful login",
    bold=True,
)

# Row heights so the block fills the page like the sample, not a short wide strip
heights = [
    0.72,
    0.72,
    0.72,
    0.72,
    0.72,
    0.72,
    1.15,
    1.15,
    0.72,
    1.35,
    1.45,
    0.85,
    0.85,
    0.85,
    1.05,
]
for i, h in enumerate(heights):
    set_row_height(table.rows[i], h)

for row in table.rows:
    for i, cell in enumerate(row.cells):
        set_cell_border(cell)
        set_cell_width(cell, COL_CM[i])

# Drop trailing empty paragraph spacing
for p in doc.paragraphs:
    if p.text != "Test report":
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

saved = []
for out in OUT_PATHS:
    try:
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out)
        saved.append(str(out))
    except OSError as exc:
        print(f"skip {out}: {exc}")

print("\n".join(saved))
